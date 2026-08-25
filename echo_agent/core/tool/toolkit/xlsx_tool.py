from pathlib import Path
from typing import Any

from langchain_core.tools import BaseTool, tool
from openpyxl import Workbook, load_workbook


def create_read_xlsx_tool() -> BaseTool:
    """创建 XLSX 读取工具"""

    @tool
    def read_xlsx(path: str, sheet_name: str | None = None) -> str:
        """
        Read data from an XLSX file.

        Use this tool when the current task requires reading spreadsheet
        content.

        If sheet_name is omitted, all sheets are read. If sheet_name is
        provided, only the specified sheet is read.

        Args:
            path: Absolute XLSX file path.
            sheet_name: Optional sheet name to read.

        Returns:
            Spreadsheet data grouped by sheet.
        """
        file_path = _resolve_xlsx_file(path)

        try:
            workbook = load_workbook(
                filename=file_path,
                read_only=True,
                data_only=True,
            )
        except OSError as exc:
            raise OSError(f"Failed to read XLSX file: {file_path}") from exc

        try:
            if sheet_name is not None:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(
                        f"Sheet not found: {sheet_name}; "
                        f"available sheets: {', '.join(workbook.sheetnames)}"
                    )
                sheet_names = [sheet_name]
            else:
                sheet_names = workbook.sheetnames

            result: list[str] = []

            for name in sheet_names:
                worksheet = workbook[name]
                result.append(f"## Sheet: {name}")

                for row in worksheet.iter_rows(values_only=True):
                    values = [_format_cell(value) for value in row]
                    result.append("\t".join(values))

            return "\n".join(result)
        finally:
            workbook.close()

    read_xlsx.metadata = {
        "annotations": {
            "readOnlyHint": True,
            "destructiveHint": False,
            "idempotentHint": True,
            "openWorldHint": False,
        }
    }
    return read_xlsx


def create_write_xlsx_tool() -> BaseTool:
    """创建 XLSX 写入工具"""

    @tool
    def write_xlsx(path: str, sheets: dict[str, list[list[Any]]]) -> str:
        """
        Write data to an XLSX file.

        Use this tool when the current task requires creating or replacing
        an XLSX spreadsheet.

        Multiple sheets can be written in a single operation.

        Args:
            path: Absolute XLSX file path.
            sheets: Mapping of sheet names to row data. Each row is a list
                of cell values.

        Returns:
            A message indicating that the XLSX file was written successfully.
        """
        if not sheets:
            raise ValueError("sheets is required")

        file_path = _resolve_xlsx_path(path)

        workbook = Workbook()

        try:
            default_sheet = workbook.active
            if default_sheet is None:
                raise RuntimeError("Failed to create default worksheet")

            workbook.remove(default_sheet)

            for sheet_name, rows in sheets.items():
                normalized_name = sheet_name.strip()
                if not normalized_name:
                    raise ValueError("Sheet name cannot be empty")

                worksheet = workbook.create_sheet(title=normalized_name)

                for row in rows:
                    worksheet.append(row)

            try:
                workbook.save(file_path)
            except OSError as exc:
                raise OSError(
                    f"Failed to write XLSX file: {file_path}"
                ) from exc
        finally:
            workbook.close()

        return (
            f"XLSX file written successfully: {file_path} "
            f"({len(sheets)} sheet(s))"
        )

    write_xlsx.metadata = {
        "annotations": {
            "readOnlyHint": False,
            "destructiveHint": True,
            "idempotentHint": True,
            "openWorldHint": False,
        }
    }
    return write_xlsx


def _resolve_xlsx_file(path: str) -> Path:
    """解析并校验 XLSX 文件路径。"""
    file_path = _resolve_xlsx_path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if not file_path.is_file():
        raise IsADirectoryError(f"Path is not a file: {file_path}")

    return file_path


def _resolve_xlsx_path(path: str) -> Path:
    """解析 XLSX 文件路径。"""
    raw = (path or "").strip()
    if not raw:
        raise ValueError("path is required")

    file_path = Path(raw).expanduser()

    if not file_path.is_absolute():
        raise ValueError(f"path must be an absolute path: {path}")

    if file_path.suffix.lower() != ".xlsx":
        raise ValueError(f"path must be an XLSX file: {path}")

    return file_path.resolve()


def _format_cell(value: Any) -> str:
    """将单元格值转换为文本。"""
    if value is None:
        return ""

    return str(value)
